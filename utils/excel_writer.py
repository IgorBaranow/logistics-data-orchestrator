import re
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .constants import CORE_INPUT_COLS, CORE_TRACKING_COLS

ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1F\x7f-\x84\x86-\x9f]")


def write_excel_dashboard(
    df: pd.DataFrame,
    final_path: str,
    columns_list: list,
    header_styles: dict,
    header_widths: dict,
    results_dict: dict,
):
    wb = Workbook()
    ws_primary = wb.worksheets[0]
    ws_primary.title = "TELEMETRY_MAIN"
    ws_audit = wb.create_sheet(title="TELEMETRY_AUDIT")

    fill_type_a = PatternFill(
        start_color="FF1B2A47", end_color="FF1B2A47", fill_type="solid"
    )
    fill_type_b = PatternFill(
        start_color="FFF0F4F8", end_color="FFF0F4F8", fill_type="solid"
    )
    fill_type_c = PatternFill(
        start_color="FFE2E8F0", end_color="FFE2E8F0", fill_type="solid"
    )

    font_type_a = Font(color="FF1B2A47", bold=True)
    font_type_b = Font(color="FF000000", bold=True)
    font_type_c = Font(color="FFFFFFFF", bold=True)
    font_link = Font(color="FF0563C1", underline="single")

    align_center = Alignment(horizontal="center", vertical="center")
    align_top = Alignment(vertical="top", wrap_text=True)
    align_right = Alignment(horizontal="right")

    for col_idx, col_name in enumerate(columns_list, 1):
        cell = ws_primary.cell(row=1, column=col_idx)

        if not isinstance(cell, Cell):
            continue

        cell.value = col_name

        if col_name in header_styles:
            cell.fill = PatternFill(
                start_color=header_styles[col_name]["fill"],
                end_color=header_styles[col_name]["fill"],
                fill_type="solid",
            )
            cell.font = Font(color=header_styles[col_name]["font"], bold=True)
        elif col_name in CORE_TRACKING_COLS:
            cell.fill, cell.font = fill_type_b, font_type_a
        elif col_name not in CORE_INPUT_COLS:
            cell.fill, cell.font = fill_type_c, font_type_b
        else:
            cell.fill, cell.font = fill_type_a, font_type_c

    for i, h in enumerate(["Entity_ID", "Detailed_Audit_Trail", "Navigation"], 1):
        c = ws_audit.cell(row=1, column=i)
        if isinstance(c, Cell):
            c.value = h
            c.fill, c.font = fill_type_a, font_type_c

    ws_audit.column_dimensions["A"].width = 20
    ws_audit.column_dimensions["B"].width = 100
    ws_audit.column_dimensions["C"].width = 20

    audit_row_idx = 2

    for row_idx, (_, row_data) in enumerate(df.iterrows(), 2):
        entity_ref = str(row_data.get("Entity_ID", "")).strip()

        for col_idx, col_name in enumerate(columns_list, 1):
            cell = ws_primary.cell(row=row_idx, column=col_idx)

            if not isinstance(cell, Cell):
                continue

            raw_val = row_data.get(col_name, "")
            val = (
                ""
                if pd.isna(raw_val)
                or str(raw_val).strip() in ["N/A", "Unknown", "None", "", "nan"]
                else raw_val
            )

            if col_name in ["Entity_ID", "Transaction_Ref", "Routing_Group_ID", "Subscriber_UID"] or (
                col_name not in CORE_TRACKING_COLS
                and col_name not in CORE_INPUT_COLS
                and "date" not in col_name.lower()
                and "timestamp" not in col_name.lower()
            ):
                if val != "":
                    val = (
                        str(int(val))
                        if isinstance(val, float) and val.is_integer()
                        else str(val)
                    )
                cell.value, cell.number_format = val, "@"

            elif col_name == "Audit_Log":
                res = results_dict.get(entity_ref)
                if res and res.get("history"):
                    h_text = ILLEGAL_XML_CHARS.sub("", str(res["history"]))

                    hist_cont_cell = ws_audit.cell(row=audit_row_idx, column=1)
                    if isinstance(hist_cont_cell, Cell):
                        hist_cont_cell.value = entity_ref
                        hist_cont_cell.font = Font(bold=True)

                    hist_val_cell = ws_audit.cell(row=audit_row_idx, column=2)
                    if isinstance(hist_val_cell, Cell):
                        hist_val_cell.value = h_text
                        hist_val_cell.alignment = align_top

                    back_cell = ws_audit.cell(row=audit_row_idx, column=3)
                    if isinstance(back_cell, Cell):
                        back_cell.value = f'=HYPERLINK("#\'TELEMETRY_MAIN\'!{cell.coordinate}", "🔙 Return")'
                        back_cell.font, back_cell.alignment = font_link, align_center

                    cell.value = f'=HYPERLINK("#\'TELEMETRY_AUDIT\'!B{audit_row_idx}", "Inspect 🔗")'
                    cell.font, cell.alignment = font_link, align_center

                    audit_row_idx += 1

            elif "date" in col_name.lower() or "eta" in col_name.lower() or "timestamp" in col_name.lower():
                if val != "" and col_name != "Next_Scheduled_State":
                    try:
                        val = (
                            datetime.strptime(val[:10], "%Y-%m-%d").date()
                            if isinstance(val, str)
                            else val.date()
                        )
                        cell.value, cell.number_format = val, "DD/MM/YYYY"
                    except Exception:
                        cell.value = val
                    if col_name in CORE_TRACKING_COLS:
                        cell.alignment = align_right
                else:
                    cell.value = val
            else:
                cell.value = val

    for i, name in enumerate(columns_list, 1):
        final_width = header_widths.get(name, 25.0)
        ws_primary.column_dimensions[get_column_letter(i)].width = final_width + 0.71

    ws_primary.auto_filter.ref = f"A1:{get_column_letter(len(columns_list))}{len(df) + 1}"
    ws_primary.freeze_panes = "B2"

    wb.save(final_path)
