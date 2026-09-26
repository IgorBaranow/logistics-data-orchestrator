import os
import shutil
import tempfile
import time
from typing import List, Tuple

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell import Cell

from .constants import CORE_INPUT_COLS, CORE_TRACKING_COLS
from .logger import get_system_logger

logger = get_system_logger("utils.remote_sync")


def extract_state_memory(
    sync_dir: str, default_manual_cols: List[str]
) -> Tuple[dict, dict, dict, List[str]]:
    historical_data, header_styles, header_widths = {}, {}, {}
    dynamic_manual_cols = default_manual_cols.copy()

    if not sync_dir or not os.path.exists(sync_dir):
        logger.info("No remote sync directory found. Using default schema.")
        return historical_data, header_styles, header_widths, dynamic_manual_cols

    try:
        existing_files = [
            f
            for f in os.listdir(sync_dir)
            if f.endswith(".xlsx") and not f.startswith("~")
        ]
    except Exception as e:
        logger.error(f"Failed to read directory {sync_dir}: {e}")
        return historical_data, header_styles, header_widths, dynamic_manual_cols

    if not existing_files:
        logger.info(f"No state files found in {sync_dir}")
        return historical_data, header_styles, header_widths, dynamic_manual_cols

    latest_file = max(
        existing_files, key=lambda x: os.path.getmtime(os.path.join(sync_dir, x))
    )
    latest_file_path = os.path.join(sync_dir, latest_file)
    logger.info(f"Extracting state memory from: {latest_file}")

    temp_snapshot = os.path.join(
        tempfile.gettempdir(), f"~temp_read_{int(time.time())}.xlsx"
    )

    try:
        shutil.copy2(latest_file_path, temp_snapshot)
    except Exception as e:
        logger.error(f"Failed to create temp snapshot of state file: {e}")
        return historical_data, header_styles, header_widths, dynamic_manual_cols

    try:
        wb_read = load_workbook(temp_snapshot, data_only=True)
        if "TELEMETRY_MAIN" in wb_read.sheetnames:
            ws_read = wb_read["TELEMETRY_MAIN"]
            for cell in ws_read[1]:
                if (
                    not isinstance(cell, Cell)
                    or not cell.value
                    or not isinstance(cell.value, str)
                ):
                    continue

                col_name = cell.value.strip()
                col_letter = cell.column_letter

                if col_letter in ws_read.column_dimensions:
                    col_dim = ws_read.column_dimensions[col_letter]
                    if col_dim.width and col_dim.width > 0:
                        header_widths[col_name] = max(col_dim.width - 0.71, 5.0)

                fill_color, font_color = None, None
                if (
                    hasattr(cell, "fill")
                    and hasattr(cell.fill, "start_color")
                    and hasattr(cell.fill.start_color, "rgb")
                ):
                    rgb = cell.fill.start_color.rgb
                    if isinstance(rgb, str) and rgb != "00000000":
                        fill_color = rgb

                if (
                    hasattr(cell, "font")
                    and hasattr(cell.font, "color")
                    and hasattr(cell.font.color, "rgb")
                ):
                    f_rgb = cell.font.color.rgb
                    if isinstance(f_rgb, str) and f_rgb != "00000000":
                        font_color = f_rgb

                if fill_color:
                    header_styles[col_name] = {
                        "fill": fill_color,
                        "font": font_color if font_color else "FF000000",
                    }
        wb_read.close()

        hist_df = pd.read_excel(temp_snapshot, sheet_name="TELEMETRY_MAIN")

        hist_df.rename(
            columns={
                "Final_Node_ETA": "Estimated_Completion",
                "ETA_Details": "Next_Scheduled_State",
            },
            inplace=True,
        )

        known_cols = set(CORE_INPUT_COLS + CORE_TRACKING_COLS)
        dynamic_manual_cols = [
            c
            for c in hist_df.columns
            if c not in known_cols and not str(c).startswith("Unnamed")
        ]

        if "Entity_ID" in hist_df.columns:
            for _, row in hist_df.iterrows():
                entity_ref = str(row["Entity_ID"]).strip().upper()
                if entity_ref and entity_ref != "NAN":
                    historical_data[entity_ref] = {
                        c: row[c]
                        for c in dynamic_manual_cols
                        if c in hist_df.columns and pd.notna(row[c])
                    }

        logger.info(f"Loaded persistence state for {len(historical_data)} entities.")

    except Exception as e:
        logger.error(f"Non-critical error reading state memory: {e}")
    finally:
        if os.path.exists(temp_snapshot):
            try:
                os.remove(temp_snapshot)
            except Exception as e:
                logger.warning(f"Could not remove temp file {temp_snapshot}: {e}")

    return historical_data, header_styles, header_widths, dynamic_manual_cols


def commit_to_remote_volume(
    local_path: str,
    sync_dir: str,
    archive_dir: str,
    file_prefix: str,
    timestamp: str,
):
    if not sync_dir or not os.path.exists(sync_dir):
        return

    logger.info("Synchronizing state to remote volume and archiving...")

    final_path = os.path.join(sync_dir, f"{file_prefix}_{timestamp}.xlsx")
    temp_sync_path = final_path + ".tmp"

    success = False
    for attempt in range(5):
        try:
            shutil.copy(local_path, temp_sync_path)
            os.replace(temp_sync_path, final_path)

            logger.info(f"SUCCESS: State file committed to remote volume -> {final_path}")
            success = True
            break
        except PermissionError:
            logger.warning(
                f"File locked by concurrent process or sync agent. Retrying... (Attempt {attempt + 1}/5)"
            )
            time.sleep(10)
        except Exception as e:
            logger.error(f"Unexpected error committing to remote volume: {e}")
            break
        finally:
            if not success and os.path.exists(temp_sync_path):
                try:
                    os.remove(temp_sync_path)
                except Exception:
                    pass

    if not success:
        logger.critical(
            f"FAILED to commit {file_prefix} after multiple attempts. Lock release timeout."
        )
        return

    try:
        if archive_dir:
            os.makedirs(archive_dir, exist_ok=True)

            all_sync_files = [
                f
                for f in os.listdir(sync_dir)
                if f.endswith(".xlsx") and not f.startswith("~")
            ]

            all_sync_files.sort(
                key=lambda x: os.path.getmtime(os.path.join(sync_dir, x))
            )

            if len(all_sync_files) > 1:
                files_to_archive = all_sync_files[:-1]
                archived_count = 0

                for f in files_to_archive:
                    src = os.path.join(sync_dir, f)
                    dst = os.path.join(archive_dir, f)

                    try:
                        shutil.copy2(src, dst)
                        delete_success = False

                        for delete_attempt in range(3):
                            try:
                                time.sleep(0.5)
                                os.remove(src)
                                delete_success = True
                                break
                            except PermissionError:
                                logger.warning(
                                    f"Entity '{f}' locked during purge. Retrying... ({delete_attempt + 1}/3)"
                                )
                                time.sleep(3)

                        if delete_success:
                            archived_count += 1
                        else:
                            logger.error(
                                f"Could not purge '{f}' post-archive. Mutex is currently held."
                            )

                    except Exception as e:
                        logger.warning(f"Could not process legacy state file '{f}'. Error: {e}")

                if archived_count > 0:
                    logger.info(
                        f"ARCHIVED: Transferred {archived_count} legacy state file(s) to vault and purged originals."
                    )
    except Exception as e:
        logger.error(f"WARNING: Remote volume archive process failed globally: {e}")
